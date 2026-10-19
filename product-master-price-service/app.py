
from datetime import datetime
from pathlib import Path
from typing import Literal
import os
import shutil
import tempfile
import stat

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from filelock import FileLock
from openpyxl import load_workbook
from pydantic import BaseModel, Field, field_validator


app = FastAPI(title="Exide Product Master Price Update Service")

WORKBOOK = Path(
    os.getenv(
        "WORKBOOK_PATH",
        "/files/Unified_Product_Price_Master_Refined_Final.xlsx",
    )
)

SHEET_NAME = "Product Master"
HEADER_ROW = 5

BACKUP_DIR = Path(
    os.getenv("BACKUP_DIR", "/files/product-master-backups")
)

PRODUCT_CODE_HEADER = "Product Code"
ALLOWED_PRICE_FIELDS = {"MRP", "Dealer Price"}

LOCK = FileLock(str(WORKBOOK) + ".update.lock")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type"],
)


class ProductPriceUpdateRequest(BaseModel):
    product_code: str = Field(min_length=1, max_length=100)
    field: Literal["MRP", "Dealer Price"]
    value: float | None = Field(
        default=None,
        ge=0,
        allow_inf_nan=False,
    )

    @field_validator("product_code")
    @classmethod
    def clean_product_code(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("Product Code is required.")
        return value


def get_headers(sheet):
    headers = {}

    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(row=HEADER_ROW, column=column).value
        if value is not None:
            headers[str(value).strip()] = column

    return headers


def find_product_row(sheet, code_column: int, product_code: str):
    target = product_code.strip()

    for row in range(HEADER_ROW + 1, sheet.max_row + 1):
        value = sheet.cell(row=row, column=code_column).value
        if value is not None and str(value).strip() == target:
            return row

    return None


def make_readable(path: Path):
    """
    Request read permissions for the workbook.

    Note: Windows bind mounts may not preserve POSIX permissions.
    The Docker/Windows host must also permit the n8n process to read it.
    """
    try:
        os.chmod(path, 0o644)
    except OSError:
        # Do not hide actual write/copy errors; this helper only handles
        # permission-setting failures. Verify readability after saving.
        pass


def verify_price(path: Path, product_code: str, field: str, expected_value):
    """Open a workbook and verify the requested price."""
    check_workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
    )

    try:
        if SHEET_NAME not in check_workbook.sheetnames:
            raise HTTPException(
                status_code=500,
                detail="Verification failed: required sheet missing.",
            )

        sheet = check_workbook[SHEET_NAME]
        headers = get_headers(sheet)

        if PRODUCT_CODE_HEADER not in headers or field not in headers:
            raise HTTPException(
                status_code=500,
                detail="Verification failed: required column missing.",
            )

        row = find_product_row(
            sheet,
            headers[PRODUCT_CODE_HEADER],
            product_code,
        )

        if row is None:
            raise HTTPException(
                status_code=500,
                detail="Verification failed: product not found.",
            )

        actual_value = sheet.cell(
            row=row,
            column=headers[field],
        ).value

        if expected_value is None:
            valid = actual_value is None or actual_value == ""
        else:
            valid = (
                actual_value is not None
                and float(actual_value) == float(expected_value)
            )

        if not valid:
            raise HTTPException(
                status_code=500,
                detail="Verification failed: saved price does not match.",
            )

        return row

    finally:
        check_workbook.close()


@app.get("/health")
def health():
    return {
        "success": True,
        "service": "exide-product-master-price-update",
        "workbook_path": str(WORKBOOK),
        "workbook_exists": WORKBOOK.is_file(),
        "sheet": SHEET_NAME,
        "header_row": HEADER_ROW,
    }


@app.post("/update-product-price")
def update_product_price(request: ProductPriceUpdateRequest):
    if request.field not in ALLOWED_PRICE_FIELDS:
        raise HTTPException(
            status_code=400,
            detail="Only MRP and Dealer Price can be updated.",
        )

    if not WORKBOOK.is_file():
        raise HTTPException(
            status_code=500,
            detail=f"Product Master workbook not found: {WORKBOOK}",
        )

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)

    with LOCK:
        workbook = None
        temp_path = None
        backup_path = None

        try:
            # Check the workbook before modifying anything.
            if not os.access(WORKBOOK, os.R_OK | os.W_OK):
                raise HTTPException(
                    status_code=500,
                    detail=(
                        "The price-update service cannot read/write the "
                        "workbook. Check Docker mount and host permissions."
                    ),
                )

            workbook = load_workbook(WORKBOOK)

            if SHEET_NAME not in workbook.sheetnames:
                raise HTTPException(
                    status_code=500,
                    detail=f"Required sheet '{SHEET_NAME}' was not found.",
                )

            sheet = workbook[SHEET_NAME]
            headers = get_headers(sheet)

            required_headers = {
                PRODUCT_CODE_HEADER,
                "MRP",
                "Dealer Price",
            }
            missing = required_headers - set(headers)

            if missing:
                raise HTTPException(
                    status_code=500,
                    detail="Missing columns: " + ", ".join(sorted(missing)),
                )

            row_number = find_product_row(
                sheet,
                headers[PRODUCT_CODE_HEADER],
                request.product_code,
            )

            if row_number is None:
                raise HTTPException(
                    status_code=404,
                    detail=(
                        f"Product Code '{request.product_code}' "
                        "was not found in Product Master."
                    ),
                )

            # Create a timestamped backup before changing the workbook.
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            backup_path = BACKUP_DIR / (
                f"{WORKBOOK.stem}_{timestamp}.xlsx"
            )
            shutil.copy2(WORKBOOK, backup_path)

            # Change only the requested price cell.
            sheet.cell(
                row=row_number,
                column=headers[request.field],
            ).value = request.value

            # Write the updated workbook to a temporary file in the
            # same directory, so it is on the same mounted filesystem.
            with tempfile.NamedTemporaryFile(
                prefix="product_master_",
                suffix=".xlsx",
                dir=str(WORKBOOK.parent),
                delete=False,
            ) as temporary_file:
                temp_path = Path(temporary_file.name)

            workbook.save(temp_path)
            workbook.close()
            workbook = None

            make_readable(temp_path)

            # Verify the temporary workbook before copying it.
            verify_price(
                temp_path,
                request.product_code,
                request.field,
                request.value,
            )

            # Copy into the existing mounted file. os.replace() was
            # avoided because it previously failed on this Windows mount.
            shutil.copyfile(temp_path, WORKBOOK)
            os.chmod(WORKBOOK, 0o644)

            # Restore readable permissions after the copy.
            make_readable(WORKBOOK)

            # Verify the final destination after the copy.
            verify_price(
                WORKBOOK,
                request.product_code,
                request.field,
                request.value,
            )

            return {
                "success": True,
                "message": "Product Master price updated and verified.",
                "product_code": request.product_code,
                "field": request.field,
                "updated_value": request.value,
                "sheet": SHEET_NAME,
                "excel_row": row_number,
                "backup_file": backup_path.name,
            }

        except HTTPException:
            raise

        except Exception as exc:
            raise HTTPException(
                status_code=500,
                detail=f"Price update failed: {type(exc).__name__}: {exc}",
            ) from exc

        finally:
            if workbook is not None:
                workbook.close()

            if temp_path is not None and temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    pass
